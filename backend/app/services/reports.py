from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.attendance import AttendanceRecord
from app.models.enums import AttendanceStatus
from app.models.grade import GradeRecord
from app.models.group import Group
from app.models.lesson_session import LessonSession
from app.models.schedule import ScheduleEntry
from app.models.student import Student

REPORT_HEADERS = ["Student", "Student #", "Attendance %", "Exam average", "Lessons marked"]


def _student_row(db: Session, *, student: Student, semester_id: int, date_from: date, date_to: date) -> list:
    attendance_stmt = (
        select(AttendanceRecord.status)
        .join(LessonSession, LessonSession.id == AttendanceRecord.session_id)
        .join(ScheduleEntry, ScheduleEntry.id == LessonSession.schedule_entry_id)
        .where(
            AttendanceRecord.student_id == student.id,
            ScheduleEntry.semester_id == semester_id,
            LessonSession.date >= date_from,
            LessonSession.date <= date_to,
        )
    )
    statuses = list(db.scalars(attendance_stmt).all())
    present = sum(1 for s in statuses if s == AttendanceStatus.PRESENT)
    attendance_pct = round(present / len(statuses) * 100, 1) if statuses else None

    grade_stmt = (
        select(GradeRecord.score)
        .join(LessonSession, LessonSession.id == GradeRecord.session_id)
        .join(ScheduleEntry, ScheduleEntry.id == LessonSession.schedule_entry_id)
        .where(
            GradeRecord.student_id == student.id,
            ScheduleEntry.semester_id == semester_id,
            LessonSession.date >= date_from,
            LessonSession.date <= date_to,
            LessonSession.is_exam.is_(True),
        )
    )
    scores = list(db.scalars(grade_stmt).all())
    exam_average = round(sum(scores) / len(scores), 1) if scores else None

    return [
        student.full_name,
        student.student_number or "",
        attendance_pct if attendance_pct is not None else "",
        exam_average if exam_average is not None else "",
        len(statuses),
    ]


def _autosize_columns(ws: Worksheet) -> None:
    for column_cells in ws.columns:
        width = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 12), 40)


def _write_group_sheet(wb: Workbook, db: Session, *, group: Group, semester_id: int, date_from: date, date_to: date) -> None:
    ws = wb.create_sheet(title=group.name[:31])
    ws.append(REPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    students = db.scalars(
        select(Student).where(Student.group_id == group.id, Student.is_active.is_(True)).order_by(Student.full_name)
    ).all()
    for student in students:
        ws.append(_student_row(db, student=student, semester_id=semester_id, date_from=date_from, date_to=date_to))

    _autosize_columns(ws)


def generate_weekly_report(
    db: Session,
    *,
    semester_id: int,
    date_from: date,
    date_to: date,
    department_id: int | None = None,
    group_id: int | None = None,
) -> bytes:
    """One worksheet per group: each active student's attendance percentage and exam-score
    average for [date_from, date_to] within the given semester."""
    groups_stmt = select(Group)
    if group_id is not None:
        groups_stmt = groups_stmt.where(Group.id == group_id)
    if department_id is not None:
        groups_stmt = groups_stmt.where(Group.department_id == department_id)
    groups = list(db.scalars(groups_stmt.order_by(Group.name)).all())

    wb = Workbook()
    wb.remove(wb.active)
    for group in groups:
        _write_group_sheet(wb, db, group=group, semester_id=semester_id, date_from=date_from, date_to=date_to)
    if not groups:
        wb.create_sheet(title="No data")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
