from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.enums import AttendanceStatus
from app.schemas.journal import AttendanceUpsert, GradeUpsert
from app.services import journal as journal_service
from app.services.reports import generate_weekly_report
from tests.factories import (
    make_assignment,
    make_department,
    make_group,
    make_room,
    make_schedule_entry,
    make_semester,
    make_student,
    make_subject,
    make_teacher,
    make_time_slot,
)


def _setup(db: Session):
    department = make_department(db)
    teacher = make_teacher(db, department)
    subject = make_subject(db, department)
    group = make_group(db, department)
    semester = make_semester(db)
    assignment = make_assignment(db, teacher, subject, group, semester)
    room = make_room(db)
    slot = make_time_slot(db)
    entry = make_schedule_entry(db, assignment, room, slot)
    return department, group, semester, entry


def test_weekly_report_has_one_sheet_per_group_with_attendance_and_exam_average(db: Session):
    department, group, semester, entry = _setup(db)
    student_a = make_student(db, group, full_name="Aisha")
    student_b = make_student(db, group, full_name="Bakyt")

    session1 = journal_service.get_or_create_session(db, schedule_entry=entry, on_date=date(2026, 9, 7))
    journal_service.set_exam_flag(session1, True)
    db.flush()
    journal_service.upsert_attendance(
        db,
        session=session1,
        records=[
            AttendanceUpsert(student_id=student_a.id, status=AttendanceStatus.PRESENT),
            AttendanceUpsert(student_id=student_b.id, status=AttendanceStatus.ABSENT),
        ],
    )
    journal_service.upsert_grades(db, session=session1, records=[GradeUpsert(student_id=student_a.id, score=80)])

    session2 = journal_service.get_or_create_session(db, schedule_entry=entry, on_date=date(2026, 9, 14))
    db.flush()
    journal_service.upsert_attendance(
        db, session=session2, records=[AttendanceUpsert(student_id=student_a.id, status=AttendanceStatus.PRESENT)]
    )
    db.flush()

    content = generate_weekly_report(
        db, semester_id=semester.id, date_from=date(2026, 9, 1), date_to=date(2026, 9, 30)
    )

    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == [group.name]
    ws = wb[group.name]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0] == ("Aisha", None, 100, 80, 2)
    assert rows[1] == ("Bakyt", None, 0, None, 1)


def test_weekly_report_filters_to_a_single_group_when_group_id_given(db: Session):
    department, group, semester, entry = _setup(db)
    other_group = make_group(db, department, name="G-999")
    make_student(db, group, full_name="Aisha")
    make_student(db, other_group, full_name="Cholpon")

    content = generate_weekly_report(
        db, semester_id=semester.id, date_from=date(2026, 9, 1), date_to=date(2026, 9, 30), group_id=group.id
    )

    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == [group.name]


def test_weekly_report_filters_by_department(db: Session):
    department, group, semester, entry = _setup(db)
    other_department = make_department(db, name="Other")

    content = generate_weekly_report(
        db,
        semester_id=semester.id,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
        department_id=other_department.id,
    )

    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["No data"]

